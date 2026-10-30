# -*- coding: utf-8 -*-
"""Büyük sayfa borç hesabı: None yedek, budget, sadece_borclu None kalır, N+1 yok.
Canlıya yazmaz. PII yok (yalnız mid / tutar / sayı).
"""
from __future__ import annotations

import os
import sys
import time
from pathlib import Path
from unittest.mock import MagicMock, patch

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
os.environ.pop("ALLOW_PROD_WRITE", None)

FAILS = []


def check(name, ok, ek=""):
    print(("ok " if ok else "FAIL ") + name + ((" :: " + str(ek)[:280]) if ek and not ok else ""))
    if not ok:
        FAILS.append(name)


def test_helpers_and_filtre():
    from routes.faturalar_routes import (
        _firma_ozet_borc_budget_sec,
        _firma_ozet_borc_fields_from_ozet,
        _firma_ozet_borc_hesaplanamadi_say,
        _firma_ozet_client_filtre_uygula,
        _firma_ozet_recalc_ozet,
        _firma_ozet_row_to_satir_item,
        _firma_ozet_sayfa_grid_hesapla,
        _firma_ozet_sum_borc_skip_none,
    )

    os.environ.pop("FIRMA_OZET_BORC_BUDGET_SEC", None)
    check("budget default ~60", 55 <= _firma_ozet_borc_budget_sec() <= 65)
    os.environ["FIRMA_OZET_BORC_BUDGET_SEC"] = "12"
    check("budget env 12", abs(_firma_ozet_borc_budget_sec() - 12) < 0.01)
    os.environ["FIRMA_OZET_BORC_BUDGET_SEC"] = "1"
    check("budget min 5", _firma_ozet_borc_budget_sec() == 5.0)
    os.environ["FIRMA_OZET_BORC_BUDGET_SEC"] = "999"
    check("budget max 110", _firma_ozet_borc_budget_sec() == 110.0)
    os.environ.pop("FIRMA_OZET_BORC_BUDGET_SEC", None)

    # Küçük sayfa: hizli + <=500 → True (prewarm yolu); büyük → False (prewarm atlanır ama borç yine hesaplanır)
    check("kucuk sayfa grid_hesapla", _firma_ozet_sayfa_grid_hesapla(True, 50) is True)
    check("buyuk sayfa prewarm kapali", _firma_ozet_sayfa_grid_hesapla(True, 5000) is False)
    check("hizli kapali her zaman", _firma_ozet_sayfa_grid_hesapla(False, 5000) is True)

    tb, ga = _firma_ozet_borc_fields_from_ozet({"toplam_borc": 42600, "geciken_ay": 3})
    check("fields ok", tb == 42600.0 and ga == 3)
    tb2, ga2 = _firma_ozet_borc_fields_from_ozet({"borc_hesaplanamadi": True, "toplam_borc": 1})
    check("fields hesaplanamadi", tb2 is None and ga2 is None)
    tb3, ga3 = _firma_ozet_borc_fields_from_ozet({})
    check("fields bos", tb3 is None and ga3 is None)

    row = {
        "id": 300,
        "firma_adi": "X",
        "current_balance": -99999,
        "musteri_durum": "aktif",
        "kira_nakit": False,
        "kdv_oran": 20,
        "firma_grid_aylik_net": 1000,
    }
    # Negatif balance yedeği YOK
    it = _firma_ozet_row_to_satir_item(row, False, None)
    check("yedek yok None", it.get("toplam_borc") is None and it.get("geciken_ay") is None)
    check("flag hesaplanamadi", it.get("borc_hesaplanamadi") is True)

    it2 = _firma_ozet_row_to_satir_item(
        row,
        False,
        {300: {"toplam_borc": 42600.0, "geciken_ay": 3, "borc_month": 14200}},
    )
    check("grid mid300", it2.get("toplam_borc") == 42600.0 and it2.get("geciken_ay") == 3)
    check("grid mid300 flag", it2.get("borc_hesaplanamadi") is False)

    it3 = _firma_ozet_row_to_satir_item(
        row, False, {300: {"borc_hesaplanamadi": True, "toplam_borc": None, "geciken_ay": None}}
    )
    check("budget none", it3.get("toplam_borc") is None and it3.get("borc_hesaplanamadi") is True)

    satirlar = [
        {"musteri_id": 1, "toplam_borc": 10.0, "aylik_tutar": 1, "firma_adi": "A", "durum_etiket": "Aktif", "hizmet_turu": "S", "liste_ara_blob": "", "grup2": ""},
        {"musteri_id": 2, "toplam_borc": None, "borc_hesaplanamadi": True, "aylik_tutar": 1, "firma_adi": "B", "durum_etiket": "Aktif", "hizmet_turu": "S", "liste_ara_blob": "", "grup2": ""},
        {"musteri_id": 3, "toplam_borc": 0.0, "aylik_tutar": 1, "firma_adi": "C", "durum_etiket": "Aktif", "hizmet_turu": "S", "liste_ara_blob": "", "grup2": ""},
        {"musteri_id": 4, "toplam_borc": 5.0, "aylik_tutar": 2, "firma_adi": "D", "durum_etiket": "Aktif", "hizmet_turu": "S", "liste_ara_blob": "", "grup2": ""},
    ]
    check("sum skip none", _firma_ozet_sum_borc_skip_none(satirlar) == 15.0)
    check("hesaplanamadi say", _firma_ozet_borc_hesaplanamadi_say(satirlar) == 1)
    ay, borc, n = _firma_ozet_recalc_ozet(satirlar)
    check("recalc", abs(ay - 5.0) < 0.01 and abs(borc - 15.0) < 0.01 and n == 4)

    filt = _firma_ozet_client_filtre_uygula(
        satirlar, q="", grup2=[], hizmet=[], durum="", oda=None, yillar=None, sadece_borclu=True
    )
    ids = [r["musteri_id"] for r in filt]
    check("sadece_borclu None kalir", 2 in ids, ids)
    check("sadece_borclu 0 gizler", 3 not in ids, ids)
    check("sadece_borclu borclu", ids == [1, 2, 4], ids)


def test_excel_none_empty():
    import io

    from openpyxl import load_workbook
    from routes.faturalar_routes import _firma_ozet_workbook_bytes

    ds = {
        "ay_referans": {"etiket": "2026-10"},
        "musteri_kapsam": "aktif",
        "total_count": 2,
        "satirlar": [
            {
                "giris_tarihi": "2024-01-01",
                "firma_adi": "A",
                "vergi_no": "—",
                "kimlik_no": "—",
                "aylik_tutar": 100,
                "ilk_kira_bedeli": None,
                "guncel_kira_bedeli": None,
                "toplam_borc": None,
                "sozlesme_gun": 1,
                "sozlesme_ay_adi": "Ocak",
                "geciken_ay": None,
                "grup2": "",
                "hizmet_turu": "S",
                "durum_etiket": "Aktif",
            },
            {
                "giris_tarihi": "2024-01-01",
                "firma_adi": "B",
                "vergi_no": "—",
                "kimlik_no": "—",
                "aylik_tutar": 100,
                "ilk_kira_bedeli": 10,
                "guncel_kira_bedeli": 10,
                "toplam_borc": 50.0,
                "sozlesme_gun": 2,
                "sozlesme_ay_adi": "Ocak",
                "geciken_ay": 3,
                "grup2": "",
                "hizmet_turu": "S",
                "durum_etiket": "Aktif",
            },
        ],
        "ozet": {"toplam_satir_kdv_dahil": 200, "toplam_borc_kdv_dahil": 50},
    }
    buf = _firma_ozet_workbook_bytes(ds)
    wb = load_workbook(buf if hasattr(buf, "read") else io.BytesIO(buf))
    ws = wb.active
    # data starts row 6
    check("excel borc None bos", ws.cell(row=6, column=8).value is None)
    check("excel ga None bos", ws.cell(row=6, column=11).value is None)
    check("excel borc dolu", ws.cell(row=7, column=8).value == 50.0)
    check("excel ga dolu", ws.cell(row=7, column=11).value == 3)


def test_batch_budget_and_cache_and_n1():
    """budget → None; canlı sonuç cache; tek toplu SELECT (N+1 yok)."""
    from datetime import date

    from routes import giris_routes as gr

    # Clear batch cache
    if hasattr(gr.musteri_firma_ozet_grid_ozet_batch, "_cache"):
        gr.musteri_firma_ozet_grid_ozet_batch._cache.clear()

    call_counts = {"payload_batch": 0, "live": 0, "fetch_all": 0}

    def fake_payload_batch(mids, *a, **k):
        call_counts["payload_batch"] += 1
        # mid 1: healthy cache; mid 2: placeholder-dominated → None (live); mid 3: no cache
        out = {}
        for mid in mids:
            if int(mid) == 1:
                out[1] = {
                    "aylar": [
                        {"yil": 2026, "ay": 10, "tutar_kdv_dahil": 1000, "brut_tutar_kdv": 1000, "odendi": False},
                        {"yil": 2026, "ay": 9, "tutar_kdv_dahil": 1000, "brut_tutar_kdv": 1000, "odendi": False},
                    ]
                }
            elif int(mid) == 2:
                out[2] = {
                    "aylar": [
                        {"yil": 2026, "ay": 10, "tutar_kdv_dahil": 0.01, "brut_tutar_kdv": 0, "odendi": False},
                        {"yil": 2026, "ay": 9, "tutar_kdv_dahil": 0.01, "brut_tutar_kdv": 0, "odendi": False},
                    ]
                }
            # mid 3: missing
        return out

    def fake_from_payload(pl, ry, rm, kyc_for_grid=None):
        if not pl:
            return None
        aylar = pl.get("aylar") or []
        # placeholder-dominated → None (canlı yol)
        ph = sum(1 for a in aylar if float(a.get("brut_tutar_kdv") or 0) <= 0)
        if ph >= len(aylar) and aylar:
            return None
        tb = sum(float(a.get("tutar_kdv_dahil") or 0) for a in aylar if not a.get("odendi"))
        return {"toplam_borc": tb, "geciken_ay": len(aylar), "borc_month": 1000.0, "borc_month_placeholder": False}

    def fake_live(mid, *a, **k):
        call_counts["live"] += 1
        time.sleep(0.05)
        return 3000.0, 3

    def fake_hucre(*a, **k):
        return 1000.0

    def fake_fetch_all(sql, params=None):
        call_counts["fetch_all"] += 1
        mids = list(params[0]) if params else []
        rows = []
        for mid in mids:
            rows.append(
                {
                    "id": mid,
                    "kyc_soz_bas": date(2024, 1, 15),
                    "giris_raw": date(2024, 1, 15),
                    "aylik_kira": 1000,
                    "kira_nakit": False,
                    "kdv_oran": 20,
                    # minimal kyc fields for _firma_ozet_kyc_dict_from_grid_sql_row
                    "mk_id": mid,
                    "sozlesme_baslangic": date(2024, 1, 15),
                }
            )
        return rows

    ref = date(2026, 10, 15)
    mids = [1, 2, 3, 4, 5]

    with patch.object(gr, "_read_aylik_grid_cache_payload_batch", side_effect=fake_payload_batch), \
         patch.object(gr, "_firma_ozet_ozet_from_grid_cache_payload", side_effect=fake_from_payload), \
         patch.object(gr, "firma_ozet_toplam_borc_ve_geciken_ay", side_effect=fake_live), \
         patch.object(gr, "firma_ozet_aylik_grid_hucre_kdv_dahil", side_effect=fake_hucre), \
         patch.object(gr, "_aylik_tahsil_edilen_aylar_set_normalized_batch", return_value={}), \
         patch.object(gr, "_plan_paket_yukle", return_value=None), \
         patch.object(gr, "prewarm_aylik_grid_cache_for_musteriler", return_value=0), \
         patch.object(gr, "_ensure_musteri_reel_donem_tutar_table", return_value=None), \
         patch.object(gr, "_tufe_map_by_year_month_cached", return_value={}), \
         patch.object(gr, "_musteri_aylik_grid_customer_kyc_select_sql", return_value="SELECT 1"), \
         patch.object(gr, "fetch_all", side_effect=fake_fetch_all), \
         patch.object(gr, "_firma_ozet_kyc_dict_from_grid_sql_row", side_effect=lambda r: {"ok": True} if r else None), \
         patch.object(gr, "_grup_rapor_ilk_ve_donem_kira", return_value=(None, None)):
        # Budget very small → after first live, rest None
        out = gr.musteri_firma_ozet_grid_ozet_batch(mids, ref, prewarm=False, budget_sec=0.08)
        meta = gr.musteri_firma_ozet_grid_ozet_batch.last_meta
        check("cache_fast mid1", (out.get(1) or {}).get("toplam_borc") == 2000.0)
        check("budget_hit", meta.get("budget_hit") is True, meta)
        n_none = sum(1 for m in mids if (out.get(m) or {}).get("borc_hesaplanamadi"))
        check("budget none sayisi>0", n_none >= 1, n_none)
        # N+1: payload_batch once per need set; fetch_all limited (customer+reel maybe)
        check("payload_batch tek", call_counts["payload_batch"] == 1, call_counts)
        check("fetch_all az", call_counts["fetch_all"] <= 4, call_counts)

        # Clear and warm: live results cached
        gr.musteri_firma_ozet_grid_ozet_batch._cache.clear()
        call_counts["live"] = 0
        out2 = gr.musteri_firma_ozet_grid_ozet_batch([3], ref, prewarm=False, budget_sec=30)
        check("canli yazdi", (out2.get(3) or {}).get("toplam_borc") == 3000.0, out2.get(3))
        check("canli sayac", call_counts["live"] == 1, call_counts)
        call_counts["live"] = 0
        out3 = gr.musteri_firma_ozet_grid_ozet_batch([3], ref, prewarm=False, budget_sec=30)
        check("sicak cache", (out3.get(3) or {}).get("toplam_borc") == 3000.0)
        check("sicak live yok", call_counts["live"] == 0, call_counts)


def test_html_ui_marks():
    html = (ROOT / "templates" / "giris" / "index.html").read_text(encoding="utf-8")
    check("ui borcFmt dash", "borcFmt = borcBilinen ? fmtF(s.toplam_borc) : '—'" in html)
    check("ui ga dash", "gaStr = gaBilinen ? String(gaNum) : '—'" in html)
    check("ui filter none keep", "boş/None: gizleme" in html or "borcRaw != null" in html)
    check("ui sum skip empty", "bRaw" in html and "if (isFinite(b)) sumBorcVis" in html)
    check("ui durum not", "borç hesaplanmadı" in html)
    check("ui meta", "borcHesaplanamadi" in html)

    src = (ROOT / "routes" / "faturalar_routes.py").read_text(encoding="utf-8")
    check("no balance fallback comment", "current_balance yedeği kaldırıldı" in src)
    check("budget env", "FIRMA_OZET_BORC_BUDGET_SEC" in src)
    check("excel ga None", '_ga_raw = s.get("geciken_ay")' in src)


def test_sadece_borclu_none_node():
    """Empty data-fr-borc-n must stay visible when sadece_borclu on."""
    import os
    import subprocess
    import tempfile

    html = (ROOT / "templates" / "giris" / "index.html").read_text(encoding="utf-8")
    a = html.find("function faturaRaporListeAraUygula")
    b = html.find("function faturaRaporListeAraTemizle", a)
    if b < 0:
        b = html.find("\nfunction ", a + 10)
    fn = html[a:b]
    js = r"""
global.document = null;
global.window = global;
function faturaRaporFirmaGrup2SatirUygun() { return true; }
function faturaRaporHizmetFiltreSatirUygun() { return true; }
function faturaRaporFirmaFiltreGrup2DiziOku() { return []; }
function faturaRaporFirmaFiltreHizmetDiziOku() { return []; }
function faturaRaporGirisAylarOku() { return []; }
function faturaRaporGirisYillarOku() { return null; }
function faturaRaporOdaFiltreOku() { return null; }
function faturaRaporFirmaListeHepsiniSecStateGuncelle() {}
function faturaRaporFirmaUpdateThFilterMarkers() {}
function faturaRaporFiltreBosSatirGuncelle() {}
function faturaRaporFirmaUstDurumGuncelle() {}
function faturaRaporFirmaSayfalamaUI() {}
function faturaRaporFiltreAktifRozetGuncelle() {}
function faturaRaporFirmaOzetAltGuncelle() {}
""" + fn + r"""
function mkTr(borc) {
  return {
    style: { display: '' },
    getAttribute: function (k) {
      if (k === 'data-fr-borc-n') return borc;
      return '';
    },
    querySelectorAll: function () { return []; }
  };
}
var rows = [mkTr('10'), mkTr(''), mkTr('0')];
var sb = { checked: true };
var gor = { value: 'firma_ozet' };
var tablo = { dataset: { faturaRaporFirma: '1' } };
var tbody = {
  querySelector: function () { return null },
  querySelectorAll: function () { return rows; }
};
global.document = {
  getElementById: function (id) {
    if (id === 'fatura_rapor_tablo') return tablo;
    if (id === 'fatura_rapor_tbody') return tbody;
    if (id === 'fatura_rapor_liste_ara') return { value: '' };
    if (id === 'fatura_rapor_gorunum') return gor;
    if (id === 'fatura_rapor_sadece_borclu') return sb;
    return null;
  }
};
faturaRaporListeAraUygula();
function say(ad, ok) { console.log((ok?'ok ':'FAIL ')+ad); if(!ok) process.exitCode=1; }
say('none gorunur', rows[1].style.display === '');
say('borclu gorunur', rows[0].style.display === '');
say('sifir gizli', rows[2].style.display === 'none');
"""
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".js", delete=False) as f:
        f.write(js)
        path = f.name
    try:
        r = subprocess.run(["node", path], capture_output=True, text=True, timeout=30)
        sys.stdout.write(r.stdout or "")
        if r.stderr:
            sys.stdout.write(r.stderr)
        check("node none exit", r.returncode == 0, r.stderr or r.stdout)
        for line in (r.stdout or "").splitlines():
            if line.startswith("FAIL "):
                FAILS.append(line[5:].split(" ::", 1)[0])
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


if __name__ == "__main__":
    test_helpers_and_filtre()
    try:
        test_excel_none_empty()
    except Exception as ex:
        check("excel suite", False, repr(ex))
    try:
        test_batch_budget_and_cache_and_n1()
    except Exception as ex:
        check("batch suite", False, repr(ex))
    test_html_ui_marks()
    test_sadece_borclu_none_node()
    if FAILS:
        print("FAILS", FAILS)
        raise SystemExit(1)
    print("PASS")
